"""生成《db-agent 新手全解》PPT。

按“Single / Multi 两个项目 × Harness 六维”组织，
并单独讲清楚数据存储、观测、安全、生命周期与数字。
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0F, 0x2A, 0x43)
BLUE = RGBColor(0x2E, 0x86, 0xDE)
TEAL = RGBColor(0x17, 0xA2, 0xB8)
ORANGE = RGBColor(0xF0, 0xA0, 0x30)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
MUTED = RGBColor(0x5A, 0x6B, 0x7B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x2B, 0x35)
SW, SH = Inches(13.333), Inches(7.5)
FONT = "PingFang SC"


def _set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def add_text(slide, x, y, w, h, lines, size=14, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for line in lines:
        item = line if isinstance(line, tuple) else (line, 0, size, bold, color)
        text, level, sz, bd, col = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.level = level
        p.line_spacing = spacing
        p.space_after = Pt(4)
        run = p.add_run()
        _set_run(run, text, sz, col, bd)
    return box


def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_header(slide, title, subtitle="", page=None):
    add_rect(slide, 0, 0, SW, Inches(0.16), fill=BLUE)
    add_text(slide, Inches(0.55), Inches(0.35), Inches(11.0), Inches(0.65),
             [(title, 0, 30, True, NAVY)])
    if subtitle:
        add_text(slide, Inches(0.57), Inches(1.02), Inches(11.0), Inches(0.4),
                 [(subtitle, 0, 14, False, MUTED)])
    if page is not None:
        add_text(slide, Inches(12.3), Inches(7.05), Inches(0.8), Inches(0.35),
                 [(str(page), 0, 12, False, MUTED)], align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, bullets, fill=LIGHT, title_color=NAVY,
             body_size=13, title_size=16):
    add_rect(slide, x, y, w, h, fill=fill)
    add_text(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), Inches(0.4),
             [(title, 0, title_size, True, title_color)])
    lines = [(("• " + b), 0, body_size, False, DARK) for b in bullets]
    add_text(slide, x + Inches(0.25), y + Inches(0.62), w - Inches(0.5), h - Inches(0.8),
             lines, size=body_size, spacing=1.12)


def add_metric_row(slide, items, y=1.55, h=1.0):
    for i, (num, label, desc) in enumerate(items):
        bx = Inches(0.7 + i * 3.0)
        add_rect(slide, bx, Inches(y), Inches(2.8), Inches(h), fill=LIGHT)
        add_text(slide, bx, Inches(y + 0.12), Inches(2.8), Inches(0.5),
                 [(num, 0, 20, True, BLUE)], align=PP_ALIGN.CENTER)
        add_text(slide, bx, Inches(y + 0.6), Inches(2.8), Inches(0.35),
                 [(label, 0, 12, True, NAVY)], align=PP_ALIGN.CENTER)
        add_text(slide, bx + Inches(0.15), Inches(y + 0.98), Inches(2.5), Inches(0.4),
                 [(desc, 0, 9, False, MUTED)], align=PP_ALIGN.CENTER)


def six_dim_slide(prs, agent_label, dim_name, bullets, page):
    s = new_slide(prs)
    add_header(s, f"{agent_label} · 六维之「{dim_name}」",
               "Harness 六维：上下文 / 记忆 / 工具 / 编排 / 观测 / 约束", page)
    add_card(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(4.6),
             dim_name, bullets, fill=LIGHT, title_color=NAVY, body_size=15, title_size=20)
    return s


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    page = 1

    # 1 封面
    s = new_slide(prs)
    add_rect(s, 0, 0, SW, SH, fill=NAVY)
    add_rect(s, 0, Inches(4.7), SW, Inches(0.06), fill=ORANGE)
    add_text(s, Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.2),
             [("db-agent 新手全解", 0, 54, True, WHITE)])
    add_text(s, Inches(1.0), Inches(2.9), Inches(11.3), Inches(0.7),
             [("Single / Multi 两个项目 × Harness 六维", 0, 24, True, TEAL)])
    add_text(s, Inches(1.0), Inches(3.7), Inches(11.3), Inches(1.0),
             [("数据怎么存 · 观测怎么做 · 安全怎么做 · 生命周期怎么走", 0, 17, False, LIGHT),
              ("给新手的操作讲解版", 0, 13, False, MUTED)])
    add_text(s, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.5),
             [("2026 · 项目讲解", 0, 13, False, MUTED)])
    page += 1

    # 2 怎么读
    s = new_slide(prs)
    add_header(s, "怎么读这份 PPT", "两条主线：Single / Multi；一把尺子：六维", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "主线一：Single Agent", [
                 "一个模型 + 15 个工具，自己决定怎么查",
                 "适合轻量、快速验证",
             ], fill=LIGHT, title_color=BLUE)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "主线二：Multi Agent", [
                 "6 个专职 Agent + Router 编排",
                 "适合复杂、多引擎、要审批的场景",
             ], fill=LIGHT, title_color=GREEN)
    add_card(s, Inches(0.7), Inches(4.45), Inches(12.0), Inches(2.1),
             "一把尺子：Harness 六维", [
                 "上下文管理 / 记忆管理 / 工具系统",
                 "系统编排 / 评估观测 / 约束与修复",
                 "每一维都会展开讲：做了什么、代码在哪、坑在哪",
             ], fill=LIGHT, title_color=NAVY)
    page += 1

    # 3 一句话 + 架构分层
    s = new_slide(prs)
    add_header(s, "项目一句话 + 架构分层", "Agent = 模型 + Harness，本项目做后者", page)
    add_text(s, Inches(0.7), Inches(1.5), Inches(12.0), Inches(0.7),
             [("一句话：让模型写 SQL 不出事——工具真执行、权限硬拦截、失败可自愈、结果可评测。",
               0, 18, True, NAVY)])
    add_card(s, Inches(0.7), Inches(2.35), Inches(3.9), Inches(1.9),
             "入口层", [
                 "CLI main.py",
                 "Streamlit app.py",
                 "Web FastAPI + SSE + React",
             ], fill=LIGHT, title_color=BLUE)
    add_card(s, Inches(4.8), Inches(2.35), Inches(3.9), Inches(1.9),
             "Harness 六维", [
                 "上下文 / 记忆 / 工具",
                 "编排 / 观测 / 约束",
                 "运行时代码在 harness/",
             ], fill=LIGHT, title_color=NAVY)
    add_card(s, Inches(8.9), Inches(2.35), Inches(3.9), Inches(1.9),
             "数据层", [
                 "SQLite 业务库",
                 "Hive / HBase 模拟",
                 "Chroma / Redis / 日志",
             ], fill=LIGHT, title_color=TEAL)
    add_text(s, Inches(0.7), Inches(4.5), Inches(12.0), Inches(1.2),
             [("Multi 代码还分了 4 层：graph（图）/ runner（执行器）/ nodes（节点）/ helpers（调度工具）。",
               0, 14, False, MUTED)])
    page += 1

    # 4 六维总览
    s = new_slide(prs)
    add_header(s, "Harness 六维总览", "六个维度，每个都回答一个新手问题", page)
    dims = [
        ("上下文", "模型看到什么？", BLUE),
        ("记忆", "模型记得什么？", TEAL),
        ("工具", "模型能做什么？", ORANGE),
        ("编排", "先做什么后做什么？", NAVY),
        ("观测", "过程怎么被记录？", GREEN),
        ("约束", "什么不能做、出错怎么办？", RED),
    ]
    for i, (t, q, c) in enumerate(dims):
        col = i % 3
        row = i // 3
        bx = Inches(0.7 + col * 4.1)
        by = Inches(1.55 + row * 2.0)
        add_rect(s, bx, by, Inches(3.85), Inches(1.75), fill=LIGHT)
        add_rect(s, bx, by, Inches(3.85), Inches(0.14), fill=c)
        add_text(s, bx + Inches(0.25), by + Inches(0.4), Inches(3.35), Inches(0.5),
                 [(t, 0, 18, True, NAVY)])
        add_text(s, bx + Inches(0.25), by + Inches(0.95), Inches(3.35), Inches(0.6),
                 [(q, 0, 13, False, MUTED)])
    page += 1

    # 5 生命周期总览（Multi 简化）
    s = new_slide(prs)
    add_header(s, "一条 Query 怎么走（Multi 总览）", "先看整体，后面会拆到每个节点", page)
    steps = [
        ("入口", "CLI/Web"),
        ("护栏", "输入检查"),
        ("Router", "意图路由"),
        ("Agent", "查数/分析"),
        ("质检", "置信度/反思"),
        ("返回", "答案+Trace"),
    ]
    x = Inches(0.55)
    for i, (t, d) in enumerate(steps):
        w = Inches(2.0)
        add_rect(s, x, Inches(2.2), w, Inches(1.2), fill=BLUE if i % 2 == 0 else NAVY)
        add_text(s, x, Inches(2.4), w, Inches(0.5),
                 [(t, 0, 16, True, WHITE)], align=PP_ALIGN.CENTER)
        add_text(s, x, Inches(2.95), w, Inches(0.4),
                 [(d, 0, 11, False, LIGHT)], align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_text(s, x + w - Inches(0.08), Inches(2.55), Inches(0.5), Inches(0.5),
                     [("→", 0, 18, True, ORANGE)], align=PP_ALIGN.CENTER)
        x += w + Inches(0.12)
    add_text(s, Inches(0.7), Inches(3.9), Inches(12.0), Inches(0.6),
             [("HITL：敏感 SQL / HBase 写 / 澄清 / 低置信度会暂停，人工批准后继续。", 0, 14, True, RED)])
    add_card(s, Inches(0.7), Inches(4.7), Inches(12.0), Inches(1.9),
             "Single 模式区别", [
                 "不做 Router / DQ / Confidence / Reflection",
                 "入口后直接进单 ReAct loop：模型自己决定调哪些工具",
             ], fill=LIGHT, title_color=BLUE)
    page += 1

    # ── Single Agent ─────────────────────────────────────────────
    s = new_slide(prs)
    add_header(s, "Single Agent 是什么", "项目一：一个模型 + 15 个工具", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "定位", [
                 "一个全能 Agent，自己决定查哪张表、调哪个工具",
                 "适合轻量查询、快速验证",
                 "入口：db-agent 默认 single 模式",
             ], fill=LIGHT, title_color=BLUE)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "关键代码", [
                 "harness/orchestration/single/agent.py：ReAct loop",
                 "tools_bundle.py：15 个工具全挂",
                 "main.py：入口 + 记忆注入",
             ], fill=LIGHT, title_color=BLUE)
    page += 1

    single_dims = [
        ("上下文管理", [
            "System Prompt 由工厂函数生成，可注入 db_type / user_role / 记忆",
            "模板优先：高频指标命中直接给预填 SQL，不调 LLM",
            "TokenBudget：估算 token，70% 水位触发压缩",
            "HybridWindowManager：最近 6 条保留原文，更早压缩成摘要",
        ]),
        ("记忆管理", [
            "短期：ConversationManager 滑动窗口 + LLM 摘要",
            "长期：VectorMemory（ChromaDB）向量召回",
            "search_memory 工具：Agent 主动查历史（Agentic RAG）",
            "自学习：成功 SQL 从 run_query 捕获，质量门 + dry-run 后回流",
        ]),
        ("工具系统", [
            "15 个工具：schema / query / analysis / chart / knowledge / hbase / hive / template / memory",
            "@tool 装饰器：函数签名自动生成 JSON Schema",
            "工具结果截断：超过 8000 字符只保留前段，避免撑爆上下文",
            "写工具幂等：save_to_memory / run_hbase 同参数 TTL 内不重复执行",
        ]),
        ("系统编排", [
            "手写 ReAct loop：Observe → Think → Act",
            "streaming 逐字输出，temperature=0 保证工具调用确定性",
            "max_turns=10 防死循环",
            "prompt caching：System Prompt + Tool Defs 做 cache_control 省 token",
        ]),
        ("评估观测", [
            "CLI：打印每轮 token 估算（ConversationManager.token_estimate）",
            "Opik：wrap_anthropic_client 包裹 client，Trace 可上报",
            "工具调用过程打印：参数、结果摘要、耗时",
            "logs/db-agent.log：LLM 响应异常与降级日志",
        ]),
        ("约束与修复", [
            "工具层 entitlement：run_query 只允许 SELECT",
            "三层护栏：输入 / SQL / 输出（下一节展开）",
            "API 重试：429 / 5xx 指数退避，未输出才重试",
            "HITL 退化：single 无 interrupt，run_query 返回「需要审批」标记",
        ]),
    ]
    for dim, bullets in single_dims:
        six_dim_slide(prs, "Single Agent", dim, bullets, page)
        page += 1

    # Single 数据存储
    s = new_slide(prs)
    add_header(s, "Single Agent · 数据怎么存", "每个文件存什么、谁读写", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "SQLite demo.db", [
                 "业务表：departments / employees / products / customers / orders",
                 "权限表：agent_roles / agent_users",
                 "用户记忆：user_memory",
                 "Feedback：user_feedback（点赞/点踩）",
             ], fill=LIGHT, title_color=NAVY)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "其它存储", [
                 "ChromaDB：长期向量记忆（harness/memory/chroma_db）",
                 "metric_registry.db：指标模板",
                 "内存：HBase 模拟、CLI 会话",
                 "logs/db-agent.log：运行日志",
             ], fill=LIGHT, title_color=NAVY)
    add_text(s, Inches(0.7), Inches(4.5), Inches(12.0), Inches(1.2),
             [("写入时机：每轮问答后写短期 + 长期记忆；成功 SQL 回流样例库（质量门通过才写）。",
               0, 14, True, NAVY)])
    page += 1

    # Single 安全
    s = new_slide(prs)
    add_header(s, "Single Agent · 安全怎么做", "三层护栏逐层展开", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "L1 输入护栏（guard_input）", [
                 "拦 prompt injection：指令覆盖 / 角色越狱",
                 "拦空输入、超长输入",
                 "在调 LLM 之前执行，零 token 成本",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "L2 SQL 护栏（run_query）", [
                 "只允许 SELECT，写操作硬拦截",
                 "权限网关：工具级 / 表级 / 行级 / 敏感列",
                 "HITL 退化为「需要审批」标记",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(0.7), Inches(4.45), Inches(12.0), Inches(2.1),
             "L3 输出护栏（guard_output）", [
                 "system prompt 泄露：硬拦截",
                 "PII（手机号/身份证/邮箱/银行卡）：只告警不拦，避免误杀正常业务结果",
                 "空输出、超长输出：硬拦截",
             ], fill=LIGHT, title_color=RED)
    page += 1

    # Single 观测
    s = new_slide(prs)
    add_header(s, "Single Agent · 观测怎么做", "轻量：token + 工具过程 + 日志", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "token 观测", [
                 "ConversationManager.token_estimate()：原文 / 摘要 / 总 token",
                 "每轮结束打印 [memory] tokens≈N",
             ], fill=LIGHT, title_color=GREEN)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "过程观测", [
                 "streaming 逐字打印回答",
                 "工具调用：名称、参数、结果摘要、耗时",
                 "重试与熔断事件打印",
             ], fill=LIGHT, title_color=GREEN)
    add_text(s, Inches(0.7), Inches(4.45), Inches(12.0), Inches(1.2),
             [("注意：single 不走 TraceContext JSONL，那是 Multi 的观测；single 靠 Opik client wrap + 打印。",
               0, 14, True, NAVY)])
    page += 1

    # Single 生命周期
    s = new_slide(prs)
    add_header(s, "Single Agent · 生命周期", "从输入到返回，按顺序走", page)
    single_steps = [
        ("输入", "用户提问"),
        ("预处理", "退出/闲聊/记忆/模板"),
        ("System", "生成 Prompt"),
        ("ReAct", "循环 ≤10 轮"),
        ("工具", "执行 + 结果回喂"),
        ("返回", "答案 + 写记忆"),
    ]
    x = Inches(0.55)
    for i, (t, d) in enumerate(single_steps):
        w = Inches(2.0)
        add_rect(s, x, Inches(2.0), w, Inches(1.2), fill=BLUE if i % 2 == 0 else NAVY)
        add_text(s, x, Inches(2.2), w, Inches(0.5),
                 [(t, 0, 15, True, WHITE)], align=PP_ALIGN.CENTER)
        add_text(s, x, Inches(2.75), w, Inches(0.4),
                 [(d, 0, 11, False, LIGHT)], align=PP_ALIGN.CENTER)
        if i < len(single_steps) - 1:
            add_text(s, x + w - Inches(0.08), Inches(2.4), Inches(0.5), Inches(0.5),
                     [("→", 0, 18, True, ORANGE)], align=PP_ALIGN.CENTER)
        x += w + Inches(0.12)
    add_card(s, Inches(0.7), Inches(3.6), Inches(12.0), Inches(2.9),
             "ReAct 循环内部（每轮）", [
                 "1. 模型决定：输出文字 或 调用工具",
                 "2. 调用工具：查 schema → 写 SQL → run_query",
                 "3. 工具结果回喂模型，模型继续判断",
                 "4. 直到不再调用工具，输出最终答案",
                 "5. 写短期 + 长期记忆，打印 token 估算",
             ], fill=LIGHT, title_color=BLUE)
    page += 1

    # Single 数字
    s = new_slide(prs)
    add_header(s, "Single Agent · 数字怎么来的", "工具 / 参数 / 测试", page)
    add_metric_row(s, [
        ("15", "工具数", "tools_bundle.py 全量挂载"),
        ("10", "最大轮数", "max_turns 防死循环"),
        ("8000", "结果截断", "超长工具结果截断字符"),
        ("300", "幂等 TTL", "写工具不重复执行(秒)"),
    ], y=1.55, h=1.5)
    add_card(s, Inches(0.7), Inches(3.4), Inches(12.0), Inches(3.1),
             "测试怎么做的（435 条中的相关部分）", [
                 "test_agent.py 19 条：工具单元 + agent_loop 集成（LLM fake）",
                 "test_guard_sql.py 23 条：SQL 护栏边界",
                 "test_rerank.py 8 条：Rerank 解析兜底",
                 "test_memory.py 30 条：短期压缩 + 向量召回",
                 "全部离线：LLM / Embedding 在测试里用脚本化 fake",
             ], fill=LIGHT, title_color=BLUE)
    page += 1

    # ── Multi Agent ───────────────────────────────────────────────
    s = new_slide(prs)
    add_header(s, "Multi Agent 是什么", "项目二：6 个专职 Agent + Router 编排", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "定位", [
                 "每个 Agent 只干一件事，Router 决定谁先谁后",
                 "适合多引擎、复杂分析、需要审批的场景",
                 "入口：db-agent --mode multi / Web / Streamlit",
             ], fill=LIGHT, title_color=GREEN)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "6 个 Agent", [
                 "SQL：只查数据",
                 "Analysis：只分析，不查库",
                 "Strategy：制度 / 指标口径",
                 "HBase：KV 操作",
                 "Hive：数仓方言",
                 "DataQuality：质量预检，只报事实",
             ], fill=LIGHT, title_color=GREEN)
    page += 1

    multi_dims = [
        ("上下文管理", [
            "Schema Linking：按查询意图检索相关表和字段（schema_discovery）",
            "值级索引：低基数 TEXT 列的真实枚举值进索引，WHERE 直接用真实值",
            "few-shot：相似问题已验证 SQL 注入（相似度阈值，命中才注入）",
            "Analysis 上下文分五层：早期摘要 / 向量记忆 / 最近对话 / Reflection 反馈 / 上游结果",
        ]),
        ("记忆管理", [
            "短期：ConversationManager 摘要 + Checkpointer 持久化 messages",
            "长期：ChromaDB / Milvus 向量记忆",
            "Self-Query：LLM 拆成语义部分 + 过滤条件（year / memory_type）",
            "自学习：run_query 成功 SQL → 质量门 → dry-run → 回流样例库",
            "HyDE + LLM Rerank：短查询扩写 + 候选精排（已接入 main.py 与 search_memory）",
        ]),
        ("工具系统", [
            "每个 Agent 按白名单挂工具，杜绝跨职责",
            "SQL Agent：discover_relevant_schema / list_tables / describe_table / run_query",
            "Analysis Agent：analyze_results / compare_periods / render_chart",
            "Strategy Agent：search_knowledge_base / lookup_metric",
            "HBase / Hive Agent：各自方言工具",
        ]),
        ("系统编排", [
            "LangGraph StateGraph：10 个节点",
            "Router：硬规则 → LRU 缓存 → LLM（确定性优先）",
            "DataQuality 首查预检（1 小时时效）",
            "Confidence Gate：SQL 自评分 < 0.7 触发 HITL",
            "Reflection ≤2 / 失败重规划 ≤1，防死循环",
            "Checkpointer：SQLite / Redis，断点续跑",
        ]),
        ("评估观测", [
            "TraceContext：每个节点一个 span，记录 node / task / elapsed / tokens / error",
            "JSONL 落盘 + Opik 双写",
            "ops_metrics：查询数 / 成功率 / token / 耗时 / 错误（Prometheus /api/metrics）",
            "告警：熔断 / Agent 超时 / 评测退化（Webhook 可选）",
            "Task board：plan 落盘 .tasks/，✓ / ✗ 可见",
        ]),
        ("约束与修复", [
            "三层护栏：输入 / SQL / 输出（展开见安全页）",
            "RBAC：5 角色，工具 / 表 / 行 / 敏感列",
            "HITL：敏感 SQL、HBase 写、澄清、低置信度，原生 interrupt",
            "API 重试 → SQL 自愈 → 重规划 → Reflection → 熔断 → 幂等",
        ]),
    ]
    for dim, bullets in multi_dims:
        six_dim_slide(prs, "Multi Agent", dim, bullets, page)
        page += 1

    # Multi 数据存储
    s = new_slide(prs)
    add_header(s, "Multi Agent · 数据怎么存", "每个存储角色：状态 / 记忆 / 审计 / 业务", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "状态类", [
                 "agent_state.db：LangGraph Checkpointer（messages / plan / results）",
                 "Redis（可选）：会话 + checkpoint，带 TTL",
                 ".tasks/*.task：Task board 落盘",
             ], fill=LIGHT, title_color=NAVY)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "记忆与业务", [
                 "demo.db：业务表 + 权限表 + user_memory + user_feedback",
                 "ChromaDB / Milvus：长期向量记忆",
                 "metric_registry.db：指标模板",
                 "uploads.db：CSV 上传",
             ], fill=LIGHT, title_color=NAVY)
    add_card(s, Inches(0.7), Inches(4.45), Inches(12.0), Inches(2.1),
             "审计与指标", [
                 "logs/traces/*.jsonl：每轮 Trace，SQL 参数脱敏",
                 "ops_metrics：进程内计数（重启清零），Prometheus 可抓取",
                 "内存：HBase 模拟、Web 会话默认内存",
             ], fill=LIGHT, title_color=NAVY)
    page += 1

    # Multi 安全
    s = new_slide(prs)
    add_header(s, "Multi Agent · 安全怎么做", "三层护栏展开 + RBAC + HITL", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "三层护栏（逐层）", [
                 "L1 输入：guard_input 拦注入 / 空 / 超长，零 token",
                 "L2 SQL：run_query 仅 SELECT，写操作硬拦",
                 "L3 输出：guard_output 拦 prompt 泄露，PII 告警",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "5 角色 RBAC", [
                 "dba / manager / analyst / viewer / support",
                 "权限存 agent_roles / agent_users，改表即生效",
                 "表级过滤 + 行级 WHERE dept_id + 敏感列",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(0.7), Inches(4.45), Inches(12.0), Inches(2.1),
             "4 类 HITL（LangGraph 原生 interrupt）", [
                 "1. 敏感列 SQL：salary / cost / budget",
                 "2. HBase 写操作：put / delete / drop / truncate",
                 "3. Clarify：模糊问题先澄清",
                 "4. Confidence Gate：SQL 低置信度请用户确认",
                 "批准后 Command(resume=...) 恢复，Checkpointer 保证状态不丢",
             ], fill=LIGHT, title_color=RED)
    page += 1

    # Multi 观测
    s = new_slide(prs)
    add_header(s, "Multi Agent · 观测怎么做", "四件套：JSONL / Opik / Prometheus / 告警", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "TraceContext JSONL", [
                 "每个节点 start_span / finish_span",
                 "记录 node / task / elapsed / tokens / error",
                 "SQL 参数脱敏后落盘 logs/traces/YYYY-MM-DD.jsonl",
             ], fill=LIGHT, title_color=GREEN)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "Opik + 指标 + 告警", [
                 "Opik：Trace / Feedback / Dataset / Experiment",
                 "ops_metrics：查询数 / 成功率 / token / 耗时 / 错误",
                 "Prometheus：/api/metrics 文本格式",
                 "告警：熔断 / Agent 超时 / 评测退化（Webhook）",
             ], fill=LIGHT, title_color=GREEN)
    add_text(s, Inches(0.7), Inches(4.5), Inches(12.0), Inches(1.2),
             [("还有：Task board 落盘执行状态；成本估算按模型统计 token（人民币 / 美元）。",
               0, 14, True, NAVY)])
    page += 1

    # Multi 生命周期
    s = new_slide(prs)
    add_header(s, "Multi Agent · 生命周期", "10 个节点按顺序走，可暂停可恢复", page)
    nodes = [
        "输入护栏", "Router", "Clarify?", "DQ?", "SQL/Hive/HBase",
        "Confidence", "Analysis", "Reflection", "Done", "收尾写记忆",
    ]
    x = Inches(0.4)
    for i, n in enumerate(nodes):
        w = Inches(1.24)
        add_rect(s, x, Inches(1.9), w, Inches(1.3), fill=NAVY if i % 2 else BLUE)
        add_text(s, x, Inches(2.15), w, Inches(0.8),
                 [(n, 0, 12, True, WHITE)], align=PP_ALIGN.CENTER)
        if i < len(nodes) - 1:
            add_text(s, x + w - Inches(0.06), Inches(2.3), Inches(0.35), Inches(0.5),
                     [("→", 0, 14, True, ORANGE)], align=PP_ALIGN.CENTER)
        x += w + Inches(0.08)
    add_card(s, Inches(0.7), Inches(3.55), Inches(12.0), Inches(3.0),
             "关键细节（数字怎么定的）", [
                 "Reflection 不通过退回 Analysis 重写，最多 2 次",
                 "Agent 超时回 Router 重规划，最多 1 次",
                 "Confidence Gate 阈值 0.7：低于则 interrupt",
                 "DataQuality 首查后 1 小时内跳过",
                 "每节点结束 Checkpointer 自动保存，HITL 暂停后 resume 接着跑",
             ], fill=LIGHT, title_color=GREEN)
    page += 1

    # Multi 数字
    s = new_slide(prs)
    add_header(s, "Multi Agent · 数字怎么来的", "435 测试 / 47 Eval / 30 Golden", page)
    add_metric_row(s, [
        ("435", "测试用例", "冒烟 62 + 单元 + 集成"),
        ("47", "Eval 用例", "guardrail 6 / routing 12 / output 16 / edge 13"),
        ("30", "Golden Set", "20 正例事实 + 10 负例"),
        ("6", "专职 Agent", "SQL / Analysis / Strategy / HBase / Hive / DQ"),
    ], y=1.55, h=1.5)
    add_card(s, Inches(0.7), Inches(3.35), Inches(12.0), Inches(3.2),
             "这些数字怎么做的", [
                 "测试按模块分层：smoke 62 / hbase 37 / entitlement 36 / orchestration 35 / bigdata 34 / memory 30",
                 "LLM-as-Judge：Kimi 评 DeepSeek，避免自评偏差",
                 "Golden Set：标准答案片段出现在回答里才算对，负例要拦截",
                 "全部离线可跑：LLM / Embedding 在测试里 fake",
             ], fill=LIGHT, title_color=GREEN)
    page += 1

    # 对比
    s = new_slide(prs)
    add_header(s, "Single vs Multi", "什么时候用哪个", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "Single", [
                 "15 个工具全挂，模型自己决定",
                 "无 Router / DQ / Confidence / Reflection",
                 "HITL 退化为「需要审批」标记",
                 "适合：轻量查询、快速验证",
             ], fill=LIGHT, title_color=BLUE)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "Multi", [
                 "6 Agent + Router + 质量闸",
                 "原生 HITL interrupt / resume",
                 "Checkpointer 持久化 + Task board",
                 "适合：多引擎、复杂分析、生产演示",
             ], fill=LIGHT, title_color=GREEN)
    add_text(s, Inches(0.7), Inches(6.3), Inches(12.0), Inches(0.6),
             [("一句话：Single 是“一个全能模型”，Multi 是“一个分工团队”。", 0, 14, True, NAVY)])
    page += 1

    # 新手误区
    s = new_slide(prs)
    add_header(s, "新手常见误区", "先把这些说清楚", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "误区 1", [
                 "“多 Agent 一定更聪明” → 不一定，分工是为了可控",
                 "“HBase/Hive 是真实集群” → 是本地模拟器，API 对齐",
             ], fill=LIGHT, title_color=ORANGE)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "误区 2", [
                 "“向量检索到处都是” → 只有记忆 / 知识检索用，SQL 不走向量",
                 "“权限靠 Prompt” → 不，工具层硬拦截才算数",
             ], fill=LIGHT, title_color=ORANGE)
    add_card(s, Inches(0.7), Inches(4.45), Inches(12.0), Inches(2.1),
             "误区 3", [
                 "“测试 435 条就是 435 个功能” → 是按模块拆的回归测试，不是功能数",
                 "“Eval 通过就是生产可用” → 当前数据源和部署仍是本地演示形态",
             ], fill=LIGHT, title_color=ORANGE)
    page += 1

    # 结尾
    s = new_slide(prs)
    add_rect(s, 0, 0, SW, SH, fill=NAVY)
    add_rect(s, 0, Inches(4.6), SW, Inches(0.06), fill=ORANGE)
    add_text(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.0),
             [("读完这份，你会讲 db-agent 了", 0, 40, True, WHITE)])
    add_text(s, Inches(1.0), Inches(4.0), Inches(11.3), Inches(0.6),
             [("Single / Multi 两条线 × 六维一把尺子", 0, 18, True, TEAL)])
    add_text(s, Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.5),
             [("数据存储 / 观测 / 安全 / 生命周期 / 数字都有出处", 0, 13, False, MUTED)])

    out = "docs/新手手册/db-agent新手全解.pptx"
    prs.save(out)
    print(f"已生成: {out}, {page} 页")


if __name__ == "__main__":
    build()
