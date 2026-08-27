"""基于 Citi 官方 PPT 模板生成 db-agent 讲解 PPT。

用法:
    python scripts/generate_project_deck_citi_template.py

输入模板:
    templates/Citi_prez_SzokePeter_template.pptx

输出:
    docs/db-agent-项目讲解-Citi模板.pptx

说明:
    - 保留模板母版 / 版式（Citi logo、页脚、页码、版式设计）
    - 清空模板原有 25 页内容，重建 20 页讲解
    - 插入本对话生成的 5 张 Mermaid 图解
"""

import struct

from pptx import Presentation
from pptx.util import Inches

TEMPLATE = "templates/Citi_prez_SzokePeter_template.pptx"
OUTPUT = "docs/db-agent-项目讲解-Citi模板.pptx"
SW, SH = 10.0, 7.5


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"不是 PNG 文件: {path}")
    w, h = struct.unpack(">II", head[16:24])
    return w, h


def delete_all_slides(prs):
    xml_slides = prs.slides._sldIdLst
    for sldId in list(xml_slides):
        xml_slides.remove(sldId)
    # 同时断开演示文稿对旧 slide part 的关系，避免保存时产生重复部件
    for rId in [r.rId for r in list(prs.part.rels.values()) if r.reltype.endswith("/slide")]:
        prs.part.drop_rel(rId)


def write_paragraphs(tf, items):
    """items: [(text, level, bold)]"""
    tf.clear()
    for i, (text, level, bold) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        for run in p.runs:
            run.font.bold = bold


def bullets(items, header=None):
    """把列表转成 (text, level, bold) 段落；header 作为一级加粗标题。"""
    out = []
    if header:
        out.append((header, 0, True))
    for b in items:
        out.append(("• " + b, 1, False))
    return out


def add_title_and_content(prs, title, body_items, layout_idx=1):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    slide.shapes.title.text_frame.text = title
    ph = slide.placeholders[10]
    write_paragraphs(ph.text_frame, body_items)
    return slide


def add_two_content(prs, title, left_header, left_items, right_header, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text_frame.text = title
    write_paragraphs(slide.placeholders[10].text_frame, bullets(left_items, left_header))
    write_paragraphs(slide.placeholders[11].text_frame, bullets(right_items, right_header))
    return slide


def add_diagram_slide(prs, title, image_path):
    slide = prs.slides.add_slide(prs.slide_layouts[3])
    slide.shapes.title.text_frame.text = title
    w_px, h_px = png_size(image_path)
    aspect = w_px / h_px
    max_w, max_h = 9.0, 4.9
    w = max_w
    h = w / aspect
    if h > max_h:
        h = max_h
        w = h * aspect
    x = Inches((SW - w) / 2)
    y = Inches(1.7 + (4.9 - h) / 2)
    slide.shapes.add_picture(image_path, x, y, width=Inches(w), height=Inches(h))
    return slide


def add_cover(prs, title, body_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.placeholders[0].text_frame.text = title
    write_paragraphs(slide.placeholders[1].text_frame, bullets(body_lines))
    return slide


def build():
    prs = Presentation(TEMPLATE)
    delete_all_slides(prs)

    # 1 封面
    add_cover(prs, "db-agent", [
        "企业级自然语言数据库分析 Agent · 工作经历项目讲解",
        "自然语言 → SQL / Hive / HBase / 指标口径 / 趋势分析",
        "独立开发 · Python / LangGraph / ChromaDB / FastAPI / React",
        "定位：Agent = 模型 + Harness",
        "本项目做 Harness：工具、权限、记忆、编排、护栏、自愈、评测",
    ])

    # 2 工作经历 → 项目背景
    add_two_content(
        prs, "工作经历 → 项目背景",
        "平时工作 · 花旗 Olympus Core（3 年 Data Analyst）", [
            "Kafka 实时流 + SFTP 批量双通道，负责全数据平台交付保障",
            "覆盖 load / validation / dispatcher / monitor / recon 全链路",
            "执行层 Spark + HDFS，调度层 Autosys",
            "四条业务线双通道稳定交付，交付质量事故 0",
            "硬指标：准时性、完整性、口径一致",
        ],
        "项目背景 · 痛点与动机", [
            "业务取数依赖专人写 SQL，沟通链路长",
            "数据源多：关系库 / 数仓 / NoSQL，语法与口径复杂",
            "Text-to-SQL 只生成不执行，不管权限、不会自愈",
            "db-agent：模型 + 工具 + 权限 + 自愈 + 评测的可交付 Harness",
        ],
    )

    # 3 项目背景与目标
    add_two_content(
        prs, "项目背景与目标",
        "业务痛点", [
            "取数依赖专人写 SQL，沟通链路长",
            "数据源多：关系库 / 数仓 / NoSQL",
            "Agent 直接查库：权限、安全、审计难控",
            "LLM 输出不稳定，容易编造数据或反复报错",
        ],
        "项目目标", [
            "自然语言自助取数，覆盖 SQL / Hive / HBase",
            "多 Agent 编排：专业分工 + 质量闸门",
            "RBAC + HITL + 护栏：越权与误操作兜底",
            "三层自愈 + 全程可观测 + 可评测",
        ],
    )

    # 4 整体架构
    add_title_and_content(prs, "整体架构", bullets([
        "入口层：CLI main.py / Streamlit app.py / Web FastAPI + SSE + React",
        "Harness 六维：编排 / 工具 / 上下文 / 记忆 / 约束 / 观测",
        "可靠性：重试 + 熔断 + 幂等 + 失败重规划 + Reflection",
        "数据层：SQLite 业务库 / Hive 模拟数仓 / HBase 内存模拟 KV",
        "状态：Checkpointer SQLite / Redis，断点续跑",
    ]))

    # 5 工程机制全景图
    add_diagram_slide(prs, "工程机制全景", "docs/diagrams/ppt/mechanism-overview-ppt.png")

    # 6 核心机制 1：多 Agent 编排
    add_title_and_content(prs, "核心机制 1：多 Agent 编排", bullets([
        "6 个专业 Agent：SQL / Analysis / Strategy / HBase / Hive / DataQuality",
        "Router 三层短路：硬规则 → LRU 缓存 → LLM",
        "质量闸：DataQuality 预检 / Confidence Gate / Reflection 重写 ≤2",
        "失败重规划 ≤1 次，防死循环",
        "Checkpointer：SQLite / Redis，HITL 断点续跑",
    ]))

    # 7 核心机制 2：安全与权限
    add_two_content(
        prs, "核心机制 2：安全与权限",
        "RBAC + 三层护栏", [
            "5 角色：dba / manager / analyst / viewer / support",
            "权限存 DB 可热改；工具 / 表 / 行三级",
            "L1 输入护栏：注入 / 空 / 超长",
            "L2 SQL 护栏：仅 SELECT",
            "L3 输出护栏：PII 告警 + prompt 泄露硬拦",
        ],
        "双层 HITL", [
            "SQL 敏感列：salary / cost / budget",
            "HBase 写操作：put / delete / drop / truncate",
            "LangGraph 原生 interrupt + resume",
            "Web API：WEB_API_TOKEN 可选鉴权",
        ],
    )

    # 8 安全与权限链路图
    add_diagram_slide(prs, "安全与权限链路", "docs/diagrams/ppt/security-chain-ppt.png")

    # 9 核心机制 3：记忆与上下文
    add_two_content(
        prs, "核心机制 3：记忆与上下文",
        "记忆系统", [
            "短期：滑动窗口 + LLM 摘要压缩",
            "长期：ChromaDB 向量召回",
            "Self-Query：先拆意图再过滤元数据",
            "TokenBudget + HybridWindow 主动压缩",
        ],
        "上下文增强", [
            "Schema Linking + 值级索引",
            "few-shot 相似 SQL 样例注入",
            "高频指标模板优先，命中零 LLM",
            "自学习：成功 SQL 工具层捕获 → dry-run → 回流",
        ],
    )

    # 10 记忆、上下文与观测闭环图
    add_diagram_slide(prs, "记忆、上下文与观测闭环", "docs/diagrams/ppt/memory-obs-loop-ppt.png")

    # 11 核心机制 4：三层自愈
    add_title_and_content(prs, "核心机制 4：三层自愈", bullets([
        "L1 API 重试：429 / 5xx / 连接错误，指数退避 + 抖动，未输出才重试",
        "L2 执行自愈：SQL 重写 ≤2；Agent 超时重规划 ≤1；Reflection 重写 ≤2",
        "L3 熔断：连续失败 ≥ 阈值 → open → 半开试探 → 恢复",
        "写工具幂等守卫：同参数 TTL 窗口内不重复执行",
    ]))

    # 12 三层自愈图
    add_diagram_slide(prs, "三层自愈", "docs/diagrams/ppt/self-healing-ppt.png")

    # 13 核心机制 5：可观测与运维
    add_title_and_content(prs, "核心机制 5：可观测与运维", bullets([
        "Trace JSONL：Span 级审计，SQL 脱敏",
        "Opik：Trace / Feedback / Dataset / Experiment",
        "Prometheus：/api/metrics（查询数 / 成功率 / token / 耗时）",
        "告警 Webhook：熔断 / Agent 超时",
        "Task board：plan 落盘 .tasks/",
        "成本估算：按模型统计 token",
    ]))

    # 14 Query 完整链路
    add_title_and_content(prs, "Query 完整链路", bullets([
        "入口 → 输入护栏 → Router → DQ（可选）→ SQL Agent",
        "SQL → 置信度门 → Analysis → Reflection → 返回",
        "HITL：敏感 SQL / HBase 写 / 澄清 / 低置信度可暂停恢复",
        "single 模式：直接 ReAct loop，无 Router / DQ / Confidence / Reflection",
    ]))

    # 15 Query 完整链路图
    add_diagram_slide(prs, "Query 完整链路", "docs/diagrams/ppt/query-flow-ppt.png")

    # 16 工程质量与成果
    add_two_content(
        prs, "工程质量与成果",
        "测试与评测", [
            "423 测试用例，离线 ~10s，零 API",
            "冒烟：模块导入 / Tool 配线 / Graph 编译",
            "Eval 35+ 用例，LLM-as-Judge（Kimi 评 DeepSeek）",
            "13 条 DB 实查事实断言",
        ],
        "产品化", [
            "CLI 安装即用 + 启动配置校验",
            "Docker Compose + 锁文件",
            "代码拆分 nodes / graph / runner / helpers",
            "类型化 state / config",
        ],
    )

    # 17 项目成果与亮点
    add_title_and_content(prs, "项目成果与亮点", bullets([
        "S：业务取数依赖专人，多数据源 + 安全合规要求高",
        "T：做一个自然语言数据库分析 Agent 的完整 Harness",
        "A：6 Agent / RBAC+HITL / 三层自愈 / 记忆闭环 / 双观测 / 评测",
        "R：3 入口覆盖 SQL / Hive / HBase；423 测试离线 10s；Eval 33/35",
        "一句话：从 Demo 做到可治理、可自愈、可评测的工程系统",
    ]))

    # 18 面试讲解要点
    add_two_content(
        prs, "面试讲解要点",
        "怎么讲", [
            "30 秒定位：Agent = 模型 + Harness",
            "2 分钟架构：入口 → 六维 → 多 Agent → 数据层",
            "深挖 1-2 个机制：优先安全或自愈",
            "收尾：评测与测试怎么保证不退化",
        ],
        "可接改进点", [
            "SQL 解析 AST 化（sqlglot）",
            "行级权限落数据库原生 RLS",
            "真实 PostgreSQL / HBase / Hive 接入",
            "Opik / Trace 打点收敛成统一 wrapper",
        ],
    )

    # 19 其他项目一览
    add_two_content(
        prs, "其他项目一览",
        "fin-agent", [
            "AI 研报分析 Agent",
            "FastAPI + LangGraph + RAG + MCP",
            "多源研报检索 / 结构化提取 / 对比分析",
            "待补充：数据源 / 编排 / 效果",
        ],
        "fraud-agent", [
            "反欺诈分析 Agent",
            "LangGraph + RAG + Streamlit + 10+ Tool",
            "双轨编排 + 知识库 + 规则引擎",
            "待补充：双轨定义 / Tool 清单 / 数据效果",
        ],
    )

    # 20 结尾
    add_cover(prs, "Thanks for your attention!", [
        "db-agent · 企业级自然语言数据库分析 Agent",
        "Q&A",
        "联系 / 演示信息占位",
    ])

    prs.save(OUTPUT)
    print(f"已生成: {OUTPUT}")


if __name__ == "__main__":
    build()
