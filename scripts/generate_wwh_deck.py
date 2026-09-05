"""生成《db-agent 四段式讲解：Who · What · When · How》PPT。"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0F, 0x2A, 0x43)
BLUE = RGBColor(0x2E, 0x86, 0xDE)
TEAL = RGBColor(0x17, 0xA2, 0xB8)
ORANGE = RGBColor(0xF0, 0xA0, 0x30)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
MUTED = RGBColor(0x5A, 0x6B, 0x7B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x2B, 0x35)
SW, SH = Inches(13.333), Inches(7.5)
FONT = "PingFang SC"


def _set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def add_text(slide, x, y, w, h, lines, size=14, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for line in lines:
        item = line if isinstance(line, tuple) else (line, 0, size, bold, color)
        text, level, sz, bd, col = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.level = level
        p.line_spacing = spacing
        p.space_after = Pt(4)
        run = p.add_run()
        _set_run(run, text, sz, col, bd)
    return box


def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_header(slide, title, subtitle="", page=None):
    add_rect(slide, 0, 0, SW, Inches(0.16), fill=BLUE)
    add_text(slide, Inches(0.55), Inches(0.35), Inches(11.0), Inches(0.65),
             [(title, 0, 30, True, NAVY)])
    if subtitle:
        add_text(slide, Inches(0.57), Inches(1.02), Inches(11.0), Inches(0.4),
                 [(subtitle, 0, 14, False, MUTED)])
    if page is not None:
        add_text(slide, Inches(12.3), Inches(7.05), Inches(0.8), Inches(0.35),
                 [(str(page), 0, 12, False, MUTED)], align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, bullets, fill=LIGHT, title_color=NAVY,
             body_size=13, title_size=16):
    add_rect(slide, x, y, w, h, fill=fill)
    add_text(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), Inches(0.4),
             [(title, 0, title_size, True, title_color)])
    lines = [(("• " + b), 0, body_size, False, DARK) for b in bullets]
    add_text(slide, x + Inches(0.25), y + Inches(0.62), w - Inches(0.5), h - Inches(0.8),
             lines, size=body_size, spacing=1.12)


def section_slide(prs, tag, title, desc, color, page):
    s = new_slide(prs)
    add_rect(s, 0, 0, SW, SH, fill=NAVY)
    add_rect(s, 0, Inches(4.6), SW, Inches(0.06), fill=color)
    add_text(s, Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.2),
             [(tag, 0, 56, True, color)])
    add_text(s, Inches(1.0), Inches(3.0), Inches(11.3), Inches(0.8),
             [(title, 0, 28, True, WHITE)])
    add_text(s, Inches(1.0), Inches(3.9), Inches(11.3), Inches(0.8),
             [(desc, 0, 16, False, LIGHT)])
    return s


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    page = 1

    # 1 封面
    s = new_slide(prs)
    add_rect(s, 0, 0, SW, SH, fill=NAVY)
    add_rect(s, 0, Inches(4.7), SW, Inches(0.06), fill=ORANGE)
    add_text(s, Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.2),
             [("db-agent 四段式讲解", 0, 52, True, WHITE)])
    add_text(s, Inches(1.0), Inches(2.9), Inches(11.3), Inches(0.7),
             [("Who · What · When · How", 0, 24, True, TEAL)])
    add_text(s, Inches(1.0), Inches(3.7), Inches(11.3), Inches(1.0),
             [("给谁用 → 是什么 → 什么时候用 → 怎么做", 0, 17, False, LIGHT)])
    add_text(s, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.5),
             [("2026 · 新手讲解版", 0, 13, False, MUTED)])
    page += 1

    # 2 框架说明
    s = new_slide(prs)
    add_header(s, "怎么用这套框架", "四段式，每段回答一个问题", page)
    items = [
        ("WHO", "给谁用、谁做的", BLUE),
        ("WHAT", "是什么、核心能力、不是什么", GREEN),
        ("WHEN", "什么时候用、什么时候不用", ORANGE),
        ("HOW", "怎么跑通、数据/安全/观测/评测怎么做", RED),
    ]
    for i, (tag, desc, c) in enumerate(items):
        bx = Inches(0.7 + i * 3.1)
        add_rect(s, bx, Inches(1.7), Inches(2.85), Inches(2.2), fill=LIGHT)
        add_rect(s, bx, Inches(1.7), Inches(2.85), Inches(0.14), fill=c)
        add_text(s, bx, Inches(2.05), Inches(2.85), Inches(0.6),
                 [(tag, 0, 24, True, c)], align=PP_ALIGN.CENTER)
        add_text(s, bx + Inches(0.2), Inches(2.8), Inches(2.45), Inches(0.9),
                 [(desc, 0, 12, False, MUTED)], align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.7), Inches(4.4), Inches(12.0), Inches(0.6),
             [("先讲 Who 和 What 建立认知，再讲 When 判断场景，最后讲 How 展示工程细节。",
               0, 14, True, NAVY)])
    page += 1

    # WHO section
    section_slide(prs, "WHO", "给谁用 · 谁做的", "先讲清楚服务对象和创作者", BLUE, page)
    page += 1

    s = new_slide(prs)
    add_header(s, "WHO · 给谁用", "四种典型用户", page)
    cards = [
        ("业务用户", "不懂 SQL，想问“上个月华东卖了多少”", BLUE),
        ("数据分析师", "想少写重复 SQL，快速拿数和口径", TEAL),
        ("开发者", "想把 NL2SQL 能力嵌进自己的系统", GREEN),
        ("管理者", "想看安全可控、可审计的数据问答", ORANGE),
    ]
    for i, (t, d, c) in enumerate(cards):
        col = i % 2
        row = i // 2
        bx = Inches(0.7 + col * 6.1)
        by = Inches(1.55 + row * 2.1)
        add_rect(s, bx, by, Inches(5.85), Inches(1.85), fill=LIGHT)
        add_text(s, bx + Inches(0.3), by + Inches(0.25), Inches(5.2), Inches(0.5),
                 [(t, 0, 18, True, NAVY)])
        add_text(s, bx + Inches(0.3), by + Inches(0.85), Inches(5.2), Inches(0.8),
                 [(d, 0, 14, False, MUTED)])
    page += 1

    s = new_slide(prs)
    add_header(s, "WHO · 谁做的", "一个人 + 数据开发背景 + 通用模型", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "我（独立开发）", [
                 "数据平台交付背景：Kafka / Spark / Hive",
                 "今年转向 Agent 应用开发",
                 "从 NL2SQL 引擎做到数据治理多 Agent",
             ], fill=LIGHT, title_color=BLUE)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "技术栈", [
                 "Python + LangGraph + FastAPI + React",
                 "DeepSeek / Anthropic 兼容 SDK",
                 "ChromaDB / Milvus / SQLite / Redis",
             ], fill=LIGHT, title_color=BLUE)
    add_text(s, Inches(0.7), Inches(6.3), Inches(12.0), Inches(0.6),
             [("定位：Agent = 模型 + Harness，我负责 Harness 那部分。", 0, 14, True, NAVY)])
    page += 1

    s = new_slide(prs)
    add_header(s, "WHO · 谁不能直接用", "诚实边界", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "需要先准备", [
                 "数据表有清晰注释 / 口径",
                 "权限模型有人维护",
                 "评测集（Golden Set）要按业务建",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "不建议直接上生产", [
                 "HBase / Hive 还是本地模拟器",
                 "SQLite 并发能力有限",
                 "外部 LLM / Embedding 有数据出境风险",
             ], fill=LIGHT, title_color=RED)
    page += 1

    # WHAT section
    section_slide(prs, "WHAT", "是什么 · 不是什么", "一句话定位 + 核心能力 + 边界", GREEN, page)
    page += 1

    s = new_slide(prs)
    add_header(s, "WHAT · 是什么", "一句话定位", page)
    add_text(s, Inches(0.7), Inches(1.7), Inches(12.0), Inches(0.8),
             [("db-agent 不是让模型写 SQL，而是让模型写 SQL 不出事。", 0, 22, True, NAVY)])
    add_text(s, Inches(0.7), Inches(2.7), Inches(12.0), Inches(0.7),
             [("工具真执行 · 权限硬拦截 · 失败可自愈 · 结果可评测", 0, 16, False, MUTED)])
    add_card(s, Inches(0.7), Inches(3.7), Inches(12.0), Inches(2.7),
             "把它拆成三层", [
                 "入口层：CLI / Streamlit / Web（FastAPI + SSE + React）",
                 "Harness 层：上下文 / 记忆 / 工具 / 编排 / 观测 / 约束",
                 "数据层：SQLite / Hive 模拟 / HBase 模拟 / 向量库 / Redis",
             ], fill=LIGHT, title_color=GREEN)
    page += 1

    s = new_slide(prs)
    add_header(s, "WHAT · 核心能力", "五个关键词", page)
    cards = [
        ("多引擎问数", "SQL / Hive / HBase", GREEN),
        ("6 Agent 编排", "查数 / 分析 / 口径 / 质量", BLUE),
        ("权限 + HITL", "RBAC + 人工审批", ORANGE),
        ("三层自愈", "重试 / 重写 / 熔断", TEAL),
        ("评测 + 自学习", "Eval + 成功 SQL 回流", NAVY),
    ]
    for i, (t, d, c) in enumerate(cards):
        col = i % 3
        row = i // 3
        bx = Inches(0.7 + col * 4.1)
        by = Inches(1.55 + row * 1.9)
        add_rect(s, bx, by, Inches(3.85), Inches(1.65), fill=LIGHT)
        add_rect(s, bx, by, Inches(3.85), Inches(0.14), fill=c)
        add_text(s, bx + Inches(0.25), by + Inches(0.35), Inches(3.35), Inches(0.5),
                 [(t, 0, 16, True, NAVY)])
        add_text(s, bx + Inches(0.25), by + Inches(0.9), Inches(3.35), Inches(0.6),
                 [(d, 0, 13, False, MUTED)])
    page += 1

    s = new_slide(prs)
    add_header(s, "WHAT · 不是什么", "防误解清单", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "它不是", [
                 "通用聊天机器人",
                 "通用 Agent 框架",
                 "真实 HBase / Hive 集群连接器",
                 "已上线的企业系统",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "它是", [
                 "数据库分析专用 Harness",
                 "可评测、可自愈、可观测的工程原型",
                 "可复用模块库：工具 / 安全 / 记忆 / 评测 / 观测",
             ], fill=LIGHT, title_color=GREEN)
    page += 1

    s = new_slide(prs)
    add_header(s, "WHAT · 六维全景", "六把尺子", page)
    dims = [
        ("上下文", "模型看到什么", BLUE),
        ("记忆", "模型记得什么", TEAL),
        ("工具", "模型能做什么", ORANGE),
        ("编排", "先做什么后做什么", NAVY),
        ("观测", "过程怎么被记录", GREEN),
        ("约束", "什么不能做、出错怎么办", RED),
    ]
    for i, (t, q, c) in enumerate(dims):
        col = i % 3
        row = i // 3
        bx = Inches(0.7 + col * 4.1)
        by = Inches(1.55 + row * 2.0)
        add_rect(s, bx, by, Inches(3.85), Inches(1.75), fill=LIGHT)
        add_rect(s, bx, by, Inches(3.85), Inches(0.14), fill=c)
        add_text(s, bx + Inches(0.25), by + Inches(0.4), Inches(3.35), Inches(0.5),
                 [(t, 0, 18, True, NAVY)])
        add_text(s, bx + Inches(0.25), by + Inches(0.95), Inches(3.35), Inches(0.6),
                 [(q, 0, 13, False, MUTED)])
    page += 1

    # WHEN section
    section_slide(prs, "WHEN", "什么时候用 · 什么时候不用", "场景判断，避免误用", ORANGE, page)
    page += 1

    s = new_slide(prs)
    add_header(s, "WHEN · 什么时候用", "四个典型场景", page)
    scenarios = [
        ("业务自助取数", "销售 / 财务 / HR 一句话查数"),
        ("数据质量预检", "查数前先扫 NULL / 日期 / 异常值"),
        ("指标口径查询", "GMV 怎么算、退款扣不扣"),
        ("趋势分析", "跨地区 / 时间对比 + 图表"),
    ]
    for i, (t, d) in enumerate(scenarios):
        col = i % 2
        row = i // 2
        bx = Inches(0.7 + col * 6.1)
        by = Inches(1.55 + row * 2.1)
        add_rect(s, bx, by, Inches(5.85), Inches(1.85), fill=LIGHT)
        add_text(s, bx + Inches(0.3), by + Inches(0.25), Inches(5.2), Inches(0.5),
                 [(t, 0, 18, True, NAVY)])
        add_text(s, bx + Inches(0.3), by + Inches(0.85), Inches(5.2), Inches(0.8),
                 [(d, 0, 14, False, MUTED)])
    page += 1

    s = new_slide(prs)
    add_header(s, "WHEN · Single 还是 Multi", "按复杂度选", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "选 Single", [
                 "轻量查询，快速验证",
                 "不想引入 Router / DQ / Reflection",
                 "只要“模型 + 15 工具”直接干活",
             ], fill=LIGHT, title_color=BLUE)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "选 Multi", [
                 "多引擎：SQL + Hive + HBase",
                 "需要 HITL 审批",
                 "需要质量闸：置信度 / 反思 / 重规划",
                 "需要 Checkpointer 断点续跑",
             ], fill=LIGHT, title_color=GREEN)
    page += 1

    s = new_slide(prs)
    add_header(s, "WHEN · 什么时候不该用", "诚实边界", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.6),
             "数据侧", [
                 "表注释 / 口径还没治理",
                 "权限模型没人维护",
                 "数据质量差到连评测集都建不出来",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.6),
             "合规侧", [
                 "敏感数据不允许发给外部 API",
                 "需要真实 HBase / Hive 集群",
                 "需要多实例高并发生产",
             ], fill=LIGHT, title_color=RED)
    page += 1

    # HOW section
    section_slide(prs, "HOW", "怎么做", "生命周期 · 数据 · 安全 · 观测 · 评测", RED, page)
    page += 1

    s = new_slide(prs)
    add_header(s, "HOW · 一条 Query 怎么走", "Multi 生命周期", page)
    nodes = [
        "输入护栏", "Router", "Clarify?", "DQ?", "SQL/Hive/HBase",
        "Confidence", "Analysis", "Reflection", "Done",
    ]
    x = Inches(0.4)
    for i, n in enumerate(nodes):
        w = Inches(1.38)
        add_rect(s, x, Inches(1.9), w, Inches(1.3), fill=NAVY if i % 2 else BLUE)
        add_text(s, x, Inches(2.15), w, Inches(0.8),
                 [(n, 0, 12, True, WHITE)], align=PP_ALIGN.CENTER)
        if i < len(nodes) - 1:
            add_text(s, x + w - Inches(0.06), Inches(2.3), Inches(0.35), Inches(0.5),
                     [("→", 0, 14, True, ORANGE)], align=PP_ALIGN.CENTER)
        x += w + Inches(0.08)
    add_card(s, Inches(0.7), Inches(3.55), Inches(12.0), Inches(3.0),
             "关键机制（数字怎么定的）", [
                 "Reflection 不通过退回重写 ≤2 次；Agent 超时重规划 ≤1 次",
                 "Confidence Gate 阈值 0.7；DQ 首查后 1 小时内跳过",
                 "每节点 Checkpointer 自动保存，HITL 暂停后 resume 接着跑",
                 "Single 模式：不走这些节点，直接 ReAct loop（≤10 轮）",
             ], fill=LIGHT, title_color=RED)
    page += 1

    s = new_slide(prs)
    add_header(s, "HOW · 数据怎么存", "每个存储的角色", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "状态与记忆", [
                 "agent_state.db：Checkpointer（messages / plan / results）",
                 "ChromaDB / Milvus：长期向量记忆",
                 "Redis（可选）：会话 + checkpoint，带 TTL",
             ], fill=LIGHT, title_color=NAVY)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "业务与审计", [
                 "demo.db：业务表 + 权限表 + user_memory + user_feedback",
                 "metric_registry.db：指标模板",
                 "logs/traces/*.jsonl：Trace 审计（SQL 脱敏）",
                 ".tasks/*.task：Task board",
             ], fill=LIGHT, title_color=NAVY)
    add_text(s, Inches(0.7), Inches(4.5), Inches(12.0), Inches(1.2),
             [("内存里还有：HBase 模拟、Web 会话默认、ops_metrics 计数（重启清零）。",
               0, 14, True, NAVY)])
    page += 1

    s = new_slide(prs)
    add_header(s, "HOW · 安全怎么做", "三层护栏逐层展开 + RBAC + HITL", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "三层护栏", [
                 "L1 输入：guard_input 拦注入 / 空 / 超长，零 token",
                 "L2 SQL：run_query 仅 SELECT，写操作硬拦",
                 "L3 输出：prompt 泄露硬拦，PII 只告警",
             ], fill=LIGHT, title_color=RED)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "权限 + 审批", [
                 "5 角色 RBAC：工具 / 表 / 行 / 敏感列",
                 "权限存 DB，改表即生效",
                 "HITL：敏感 SQL / HBase 写 / 澄清 / 低置信度",
                 "Web API 可选 Bearer Token",
             ], fill=LIGHT, title_color=RED)
    add_text(s, Inches(0.7), Inches(4.5), Inches(12.0), Inches(1.2),
             [("安全结论：不靠模型自觉，靠工具层硬拦截 + 人工审批 + 审计。",
               0, 14, True, NAVY)])
    page += 1

    s = new_slide(prs)
    add_header(s, "HOW · 观测怎么做", "四件套 + Task board", page)
    add_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(2.7),
             "Trace / Opik", [
                 "TraceContext：每节点 span（node / task / elapsed / tokens / error）",
                 "JSONL 落盘 + Opik Trace / Feedback",
             ], fill=LIGHT, title_color=GREEN)
    add_card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.7),
             "指标 / 告警", [
                 "ops_metrics：查询数 / 成功率 / token / 耗时 / 错误",
                 "Prometheus /api/metrics",
                 "告警：熔断 / 超时 / 评测退化（Webhook）",
             ], fill=LIGHT, title_color=GREEN)
    add_text(s, Inches(0.7), Inches(4.5), Inches(12.0), Inches(1.2),
             [("还有 Task board：plan 落盘 .tasks/，每一步 ✓ / ✗ 可见。",
               0, 14, True, NAVY)])
    page += 1

    s = new_slide(prs)
    add_header(s, "HOW · 评测怎么做", "数字都有出处", page)
    metrics = [
        ("435", "测试用例", "smoke 62 + 单元 + 集成"),
        ("47", "Eval", "guardrail 6 / routing 12 / output 16 / edge 13"),
        ("30", "Golden", "20 正例事实 + 10 负例"),
        ("15", "工具", "Single 全量挂载"),
    ]
    for i, (num, label, desc) in enumerate(metrics):
        bx = Inches(0.7 + i * 3.0)
        add_rect(s, bx, Inches(1.6), Inches(2.8), Inches(1.5), fill=LIGHT)
        add_text(s, bx, Inches(1.72), Inches(2.8), Inches(0.5),
                 [(num, 0, 22, True, BLUE)], align=PP_ALIGN.CENTER)
        add_text(s, bx, Inches(2.25), Inches(2.8), Inches(0.35),
                 [(label, 0, 12, True, NAVY)], align=PP_ALIGN.CENTER)
        add_text(s, bx + Inches(0.15), Inches(2.6), Inches(2.5), Inches(0.4),
                 [(desc, 0, 9, False, MUTED)], align=PP_ALIGN.CENTER)
    add_card(s, Inches(0.7), Inches(3.45), Inches(12.0), Inches(3.1),
             "怎么保证不是自说自话", [
                 "LLM-as-Judge：Kimi 评 DeepSeek，避免自评偏差",
                 "Golden Set：标准答案片段出现在回答里才算对",
                 "全部离线可跑：LLM / Embedding 在测试里 fake",
                 "检索消融评测：量化 裸向量 / +HyDE / +Rerank 各自贡献",
             ], fill=LIGHT, title_color=GREEN)
    page += 1

    # 结尾
    s = new_slide(prs)
    add_rect(s, 0, 0, SW, SH, fill=NAVY)
    add_rect(s, 0, Inches(4.6), SW, Inches(0.06), fill=ORANGE)
    add_text(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.0),
             [("Who · What · When · How", 0, 44, True, WHITE)])
    add_text(s, Inches(1.0), Inches(4.0), Inches(11.3), Inches(0.6),
             [("给谁用 → 是什么 → 什么时候用 → 怎么做", 0, 18, True, TEAL)])
    add_text(s, Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.5),
             [("db-agent 新手讲解 · 四段式", 0, 13, False, MUTED)])

    out = "docs/项目介绍/db-agent-WWH讲解.pptx"
    prs.save(out)
    print(f"已生成: {out}, {page} 页")


if __name__ == "__main__":
    build()
